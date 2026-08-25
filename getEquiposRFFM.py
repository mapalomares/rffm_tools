#!/usr/bin/env python3
"""
getEquiposRFFM.py - v1.0

Autor: Miguel Ángel Palomares
Fecha: 2026-08-25
Version: 1.0

Descripcion:
    Dada la URL de la clasificacion de un grupo de rffm.es, recorre cada
    equipo (siguiendo el enlace a su ficha) y recoge Equipo, Localidad,
    Terreno de juego y Equipacion (camiseta/pantalon/medias). El resultado
    se exporta a Excel.

Parametros de entrada:
    - url (posicional u opcional --url): URL de la clasificacion del grupo,
      p.ej. https://www.rffm.es/competicion/clasificaciones?temporada=22&tipojuego=1&competicion=26737828&grupo=26737840&jornada=1
      El parametro 'jornada' es irrelevante y puede omitirse.
    - --conf: ruta del fichero .conf (por defecto getEquiposRFFM.conf)

Parametros de salida:
    - output/equipos_<competicion>_<grupo>_<temporada>.xlsx
    - logs/getEquiposRFFM_<timestamp>.log

Uso: python getEquiposRFFM.py "<url de la clasificacion>"
"""

import argparse
import configparser
import json
import logging
import os
import re
import sys
import time
from datetime import datetime

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONF = os.path.join(SCRIPT_DIR, "getEquiposRFFM.conf")
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) getEquiposRFFM/1.0"
HEADERS = ["Equipo", "Localidad", "Terreno de juego", "Equipacion"]
HEADER_COLOR = "FF305496"
MINOR_WORDS = {"de", "del", "la", "las", "el", "los", "y", "en"}

log = logging.getLogger("getEquiposRFFM")


def load_config(conf_path):
    if not os.path.isfile(conf_path):
        print(f"[ERROR] Fichero de configuracion no encontrado: {conf_path}")
        sys.exit(1)
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(conf_path, encoding="utf-8")
    return cfg


def setup_logging(cfg):
    folder = resolve_path(cfg.get("logs", "folder", fallback="logs"))
    os.makedirs(folder, exist_ok=True)
    level = getattr(logging, cfg.get("logs", "level", fallback="INFO").upper(), logging.INFO)
    logfile = os.path.join(folder, f"getEquiposRFFM_{datetime.now():%Y%m%d_%H%M%S}.log")
    logging.basicConfig(
        level=level,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.StreamHandler(sys.stdout),
            logging.FileHandler(logfile, encoding="utf-8"),
        ],
    )
    log.info(f"Log: {logfile}")


def resolve_path(path):
    return path if os.path.isabs(path) else os.path.join(SCRIPT_DIR, path)


def fetch_next_data(url, timeout):
    """Descarga una pagina de rffm.es y devuelve el bloque JSON __NEXT_DATA__."""
    resp = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=timeout)
    resp.raise_for_status()
    match = re.search(
        r'<script id="__NEXT_DATA__" type="application/json">(.*?)</script>',
        resp.text,
        re.DOTALL,
    )
    if not match:
        log.error(f"No se ha encontrado el bloque __NEXT_DATA__ en: {url}")
        sys.exit(2)
    return json.loads(match.group(1))


def fetch_standings(url, timeout):
    """Descarga la clasificacion y devuelve (info_grupo, lista_equipos)."""
    log.info(f"Descargando clasificacion: {url}")
    data = fetch_next_data(url, timeout)
    standings = data.get("props", {}).get("pageProps", {}).get("standings")
    if not standings or not standings.get("clasificacion"):
        log.error("La respuesta no contiene clasificacion (revisa los parametros de la URL).")
        sys.exit(2)
    season = data.get("props", {}).get("pageProps", {}).get("season") or {}
    info = {
        "competicion": standings.get("competicion", "competicion"),
        "grupo": standings.get("grupo", "grupo"),
        "temporada": season.get("nombre", "") if isinstance(season, dict) else season,
    }
    equipos = [
        {"codequipo": e.get("codequipo"), "nombre": e.get("nombre")}
        for e in standings["clasificacion"]
    ]
    log.info(f"Equipos en la clasificacion: {len(equipos)}")
    return info, equipos


def spanish_title(text):
    """Title-case que mantiene en minuscula las preposiciones/articulos internos."""
    words = text.lower().split()
    result = [
        w if i > 0 and w in MINOR_WORDS else w.capitalize() for i, w in enumerate(words)
    ]
    return " ".join(result)


def clean_team_name(nombre):
    """'A.D. EL NORTE 'A'' -> 'A.D. El Norte A'."""
    sin_comillas = nombre.replace("'", "").replace("\u2019", "")
    return sin_comillas.title()


def split_campo(campo):
    """'GETAFE - JUAN DE LA CIERVA 2 (HA)' -> 'Juan de la Cierva 2 (Getafe)'.

    Primero se elimina la anotacion final entre parentesis (superficie, etc.),
    que puede contener guiones, y despues se separa el prefijo de localidad.
    """
    if not campo:
        return ""
    sin_anotacion = re.sub(r"\s*\([^()]*\)\s*$", "", campo).strip()
    partes = sin_anotacion.split(" - ", 1)
    if len(partes) == 2:
        localidad_campo, nombre_campo = partes
    else:
        localidad_campo, nombre_campo = "", partes[0]
    nombre_campo = spanish_title(nombre_campo.strip())
    if localidad_campo.strip():
        return f"{nombre_campo} ({spanish_title(localidad_campo.strip())})"
    return nombre_campo


def fetch_team_details(codequipo, timeout):
    """Descarga la ficha de un equipo y devuelve equipo/localidad/terreno/equipacion."""
    url = f"https://www.rffm.es/fichaequipo/{codequipo}"
    data = fetch_next_data(url, timeout)
    team = data.get("props", {}).get("pageProps", {}).get("team")
    if not team:
        log.warning(f"Sin datos de ficha para el equipo {codequipo}: {url}")
        return None

    equipaciones = team.get("equipaciones") or []
    principal = equipaciones[0] if equipaciones else {}
    equipacion = "/".join(
        (principal.get(campo) or "").title() for campo in ("camiseta", "pantalon", "medias")
    )

    return {
        "equipo": clean_team_name(team.get("nombre_equipo", "")),
        "localidad": team.get("localidad_correspondencia", "").strip(),
        "terreno": split_campo(team.get("campo", "")),
        "equipacion": equipacion,
    }


def collect_teams(equipos, timeout, delay):
    """Recorre cada equipo de la clasificacion y descarga su ficha."""
    rows = []
    for i, equipo in enumerate(equipos, start=1):
        codequipo = equipo["codequipo"]
        log.info(f"[{i}/{len(equipos)}] {equipo['nombre']} (equipo {codequipo})")
        details = fetch_team_details(codequipo, timeout)
        if details:
            rows.append(details)
        if delay and i < len(equipos):
            time.sleep(delay)
    return rows


def slugify(text):
    text = re.sub(r"[^\w\s-]", "", str(text), flags=re.UNICODE).strip()
    return re.sub(r"[\s]+", "_", text)


def build_path(cfg, info):
    folder = resolve_path(cfg.get("output", "folder", fallback="output"))
    os.makedirs(folder, exist_ok=True)
    filename = cfg.get("output", "filename").format(
        competicion=slugify(info.get("competicion", "competicion")),
        grupo=slugify(info.get("grupo", "grupo")),
        temporada=slugify(info.get("temporada", "")),
    )
    return os.path.join(folder, filename)


def write_excel(rows, info, cfg):
    path = build_path(cfg, info)

    wb = Workbook()
    ws = wb.active
    ws.title = cfg.get("output", "sheet_name", fallback="Equipos")

    ws.append(HEADERS)
    header_fill = PatternFill("solid", fgColor=HEADER_COLOR)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row["equipo"], row["localidad"], row["terreno"], row["equipacion"]])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    for col, width in enumerate([32, 18, 34, 34], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(path)
    log.info(f"Excel generado: {path}")
    return path


def main():
    parser = argparse.ArgumentParser(
        description="Recoge Equipo/Localidad/Terreno/Equipacion de todos los equipos de un grupo de RFFM."
    )
    parser.add_argument(
        "url", nargs="?", help="URL completa de la clasificacion (si se omite, se usa la del .conf)"
    )
    parser.add_argument("--conf", default=DEFAULT_CONF, help="Ruta del fichero .conf")
    parser.add_argument("--url", dest="url_opt", help="Alternativa a la URL posicional")
    args = parser.parse_args()

    cfg = load_config(args.conf)
    setup_logging(cfg)

    url = args.url or args.url_opt
    if url:
        cfg.set("source", "url", url)

    timeout = cfg.getint("source", "timeout", fallback=30)
    delay = cfg.getfloat("source", "delay", fallback=0.5)

    info, equipos = fetch_standings(cfg.get("source", "url"), timeout)
    log.info(f"Competicion: {info['competicion']} | {info['grupo']} | Temporada {info['temporada']}")

    rows = collect_teams(equipos, timeout, delay)
    write_excel(rows, info, cfg)


if __name__ == "__main__":
    try:
        main()
    except requests.RequestException as exc:
        log.error(f"Error de red: {exc}")
        sys.exit(3)
