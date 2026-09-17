#!/usr/bin/env python3
"""
getCalendarioRFFM.py - v1.2

Descarga el calendario de una competición de rffm.es y lo exporta a Excel y PDF
con las columnas: Jornada, Fecha, Local, Visitante.
Las filas del equipo elegido se marcan con color y cada jornada se separa con
una línea doble.

Uso: python getCalendarioRFFM.py "<url del calendario>"
"""

import argparse
import configparser
import json
import logging
import os
import re
import sys
from datetime import datetime
from functools import partial

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONF = os.path.join(SCRIPT_DIR, "getCalendarioRFFM.conf")
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) getCalendarioRFFM/1.0"
HEADERS = ["Jornada", "Fecha", "Local", "Visitante"]
HEADER_COLOR = "FF305496"

log = logging.getLogger("getCalendarioRFFM")


def load_config(conf_path):
    if not os.path.isfile(conf_path):
        print(f"[ERROR] Fichero de configuración no encontrado: {conf_path}")
        sys.exit(1)
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(conf_path, encoding="utf-8")
    return cfg


def setup_logging(cfg):
    folder = resolve_path(cfg.get("logs", "folder", fallback="logs"))
    os.makedirs(folder, exist_ok=True)
    level = getattr(logging, cfg.get("logs", "level", fallback="INFO").upper(), logging.INFO)
    logfile = os.path.join(folder, f"getCalendarioRFFM_{datetime.now():%Y%m%d_%H%M%S}.log")
    logging.basicConfig(
        level=level,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.StreamHandler(sys.stdout),
            logging.FileHandler(logfile, encoding="utf-8"),
        ],
    )
    log.info(f"Log: {logfile}")


def resolve_path(path):
    return path if os.path.isabs(path) else os.path.join(SCRIPT_DIR, path)


def fetch_calendar(url, timeout, verify=True):
    """Descarga la página y extrae el bloque JSON __NEXT_DATA__ con el calendario."""
    log.info(f"Descargando: {url}")
    resp = requests.get(
        url,
        headers={"User-Agent": USER_AGENT},
        timeout=timeout,
        verify=verify,
    )
    resp.raise_for_status()

    match = re.search(
        r'<script id="__NEXT_DATA__" type="application/json">(.*?)</script>',
        resp.text,
        re.DOTALL,
    )
    if not match:
        log.error("No se ha encontrado el bloque __NEXT_DATA__ en la página.")
        sys.exit(2)

    data = json.loads(match.group(1))
    calendar = data.get("props", {}).get("pageProps", {}).get("calendar")
    if not calendar or not calendar.get("rounds"):
        log.error("La respuesta no contiene calendario (revisa los parámetros de la URL).")
        sys.exit(2)
    return calendar


def extract_rows(calendar):
    """Devuelve una lista de dicts con jornada, fecha, local y visitante."""
    rows = []
    for rnd in calendar.get("rounds", []):
        jornada = rnd.get("codjornada") or rnd.get("jornada", "")
        for match in rnd.get("equipos", []):
            rows.append(
                {
                    "jornada": int(jornada) if str(jornada).isdigit() else jornada,
                    "fecha": parse_date(match.get("fecha", "")),
                    "local": match.get("equipo_local", "").strip(),
                    "visitante": match.get("equipo_visitante", "").strip(),
                }
            )
    log.info(f"Partidos extraídos: {len(rows)} en {len(calendar.get('rounds', []))} jornadas")
    return rows


def list_teams(rows):
    """Lista ordenada de equipos distintos que aparecen en el calendario."""
    return sorted({r["local"] for r in rows} | {r["visitante"] for r in rows})


def choose_team(teams, default=""):
    """Pregunta por consola qué equipo resaltar. Devuelve '' si no se resalta ninguno."""
    if not sys.stdin.isatty():
        log.info("Entrada no interactiva: se mantiene el equipo configurado.")
        return default

    default_idx = next(
        (i for i, t in enumerate(teams, start=1) if t.upper() == default.strip().upper()), None
    )

    print("\nEquipos del grupo:")
    print("   0. (ninguno, no resaltar)")
    for i, team in enumerate(teams, start=1):
        mark = " *" if i == default_idx else ""
        print(f"  {i:2}. {team}{mark}")

    suffix = f" [{default_idx}]" if default_idx else " [0]"
    while True:
        answer = input(f"\nEquipo a resaltar{suffix}: ").strip()
        if not answer:
            return teams[default_idx - 1] if default_idx else ""
        if answer.isdigit() and 0 <= int(answer) <= len(teams):
            idx = int(answer)
            return teams[idx - 1] if idx else ""
        matches = [t for t in teams if answer.upper() in t.upper()]
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            print(f"Hay {len(matches)} coincidencias, concreta más o usa el número.")
        else:
            print("Opción no válida.")


def parse_date(value):
    """Convierte 'dd-mm-yyyy' en date; si no es válida, devuelve el texto original."""
    try:
        return datetime.strptime(value.strip(), "%d-%m-%Y").date()
    except (ValueError, AttributeError):
        return value


def slugify(text):
    text = re.sub(r"[^\w\s-]", "", str(text), flags=re.UNICODE).strip()
    return re.sub(r"[\s]+", "_", text)


def build_path(cfg, calendar, key):
    """Ruta de salida a partir de la plantilla de nombre indicada en [output]."""
    folder = resolve_path(cfg.get("output", "folder", fallback="output"))
    os.makedirs(folder, exist_ok=True)
    filename = cfg.get("output", key).format(
        competicion=slugify(calendar.get("competicion", "competicion")),
        grupo=slugify(calendar.get("grupo", "grupo")),
        temporada=slugify(calendar.get("temporada", "")),
    )
    return os.path.join(folder, filename)


def is_highlighted(row, team_upper):
    return bool(team_upper) and (
        team_upper in row["local"].upper() or team_upper in row["visitante"].upper()
    )


def is_round_end(rows, idx):
    return idx + 1 == len(rows) or rows[idx + 1]["jornada"] != rows[idx]["jornada"]


def write_excel(rows, calendar, cfg):
    path = build_path(cfg, calendar, "filename")

    highlight_team = cfg.get("highlight", "team", fallback="").strip().upper()
    fill = PatternFill("solid", fgColor=cfg.get("highlight", "color", fallback="FFFFF2CC"))

    wb = Workbook()
    ws = wb.active
    ws.title = cfg.get("output", "sheet_name", fallback="Calendario")

    headers = HEADERS
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor=HEADER_COLOR)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    round_border = Border(bottom=Side(style="double", color="FF000000"))

    marked = 0
    for idx, row in enumerate(rows):
        ws.append([row["jornada"], row["fecha"], row["local"], row["visitante"]])
        r = ws.max_row
        ws.cell(row=r, column=2).number_format = "DD/MM/YYYY"
        if is_highlighted(row, highlight_team):
            marked += 1
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).fill = fill
                ws.cell(row=r, column=col).font = Font(bold=True)

        if is_round_end(rows, idx):
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).border = round_border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    for col, width in enumerate([10, 14, 42, 42], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(path)
    log.info(f"Filas resaltadas para '{highlight_team or '(ninguno)'}': {marked}")
    log.info(f"Excel generado: {path}")
    return path


def argb_color(value, default):
    """Convierte un color ARGB/RGB hexadecimal del .conf en un color de reportlab."""
    rgb = (value or default).strip().lstrip("#")
    if len(rgb) == 8:
        rgb = rgb[2:]
    return colors.HexColor(f"#{rgb}")


def watermark_image(cfg, team):
    """Ruta de la marca de agua si el equipo elegido coincide con el patrón del .conf."""
    if not cfg.getboolean("watermark", "enabled", fallback=False) or not team:
        return None
    pattern = cfg.get("watermark", "team_match", fallback="").strip().upper()
    if not pattern or pattern not in team.upper():
        return None
    image = resolve_path(cfg.get("watermark", "image", fallback=""))
    if not os.path.isfile(image):
        log.warning(f"Marca de agua no encontrada: {image}")
        return None
    return image


def tile_watermark(canvas, doc, image, opacity, table_heights):
    """Repite la imagen en vertical, del ancho de la tabla y solo mientras haya tabla."""
    height = table_heights[doc.page - 1] if doc.page <= len(table_heights) else 0
    if height <= 0:
        return

    img = ImageReader(image)
    iw, ih = img.getSize()
    width = doc.width
    tile_h = width * ih / iw
    top = doc.bottomMargin + doc.height

    canvas.saveState()
    clip = canvas.beginPath()
    clip.rect(doc.leftMargin, top - height, width, height)
    canvas.clipPath(clip, stroke=0)
    canvas.setFillAlpha(opacity)
    y = top - tile_h
    while y > top - height - tile_h:
        canvas.drawImage(img, doc.leftMargin, y, width, tile_h, mask="auto")
        y -= tile_h
    canvas.restoreState()


def split_heights(table, avail_w, avail_h):
    """Altura ocupada por la tabla en cada página."""
    heights = []
    rest = table
    for _ in range(1000):
        parts = rest.split(avail_w, avail_h)
        if len(parts) < 2:
            heights.append(rest.wrap(avail_w, avail_h)[1])
            break
        heights.append(parts[0].wrap(avail_w, avail_h)[1])
        rest = parts[1]
    return heights


def write_pdf(
    rows, calendar, cfg, filename_key="pdf_filename", label="PDF", highlight_rows=True
):
    path = build_path(cfg, calendar, filename_key)
    highlight_team = cfg.get("highlight", "team", fallback="").strip().upper()
    highlight_color = argb_color(cfg.get("highlight", "color", fallback=""), "FFFFF2CC")

    data = [HEADERS]
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), argb_color(HEADER_COLOR, HEADER_COLOR)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 1), (1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
    ]

    for idx, row in enumerate(rows):
        fecha = row["fecha"]
        data.append(
            [
                str(row["jornada"]),
                fecha.strftime("%d/%m/%Y") if hasattr(fecha, "strftime") else str(fecha),
                row["local"],
                row["visitante"],
            ]
        )
        r = idx + 1
        if highlight_rows and is_highlighted(row, highlight_team):
            style.append(("BACKGROUND", (0, r), (-1, r), highlight_color))
            style.append(("FONTNAME", (0, r), (-1, r), "Helvetica-Bold"))
        if is_round_end(rows, idx):
            # count=2 dibuja la línea doble de separación entre jornadas
            style.append(("LINEBELOW", (0, r), (-1, r), 0.6, colors.black, None, None, None, 2, 1.5))

    doc = SimpleDocTemplate(
        path,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"{calendar.get('competicion')} - {calendar.get('grupo')}",
    )
    table = Table(data, colWidths=[17 * mm, 23 * mm, 70 * mm, 70 * mm], repeatRows=1)
    table.setStyle(TableStyle(style))

    image = watermark_image(cfg, cfg.get("highlight", "team", fallback=""))
    if image:
        opacity = cfg.getfloat("watermark", "opacity", fallback=0.12)
        heights = split_heights(table, doc.width, doc.height)
        on_page = partial(tile_watermark, image=image, opacity=opacity, table_heights=heights)
        log.info(f"Marca de agua aplicada: {image}")
        doc.build([table], onFirstPage=on_page, onLaterPages=on_page)
    else:
        doc.build([table])

    log.info(f"{label} generado: {path}")
    return path


def main():
    parser = argparse.ArgumentParser(description="Exporta el calendario RFFM a Excel y PDF.")
    parser.add_argument(
        "url", nargs="?", help="URL completa del calendario (si se omite, se usa la del .conf)"
    )
    parser.add_argument("--conf", default=DEFAULT_CONF, help="Ruta del fichero .conf")
    parser.add_argument("--url", dest="url_opt", help="Alternativa a la URL posicional")
    parser.add_argument("--team", help="Equipo a resaltar; omite la pregunta interactiva")
    parser.add_argument(
        "--no-prompt", action="store_true", help="No preguntar: usa el equipo del .conf"
    )
    parser.add_argument("--no-pdf", action="store_true", help="No generar el PDF")
    args = parser.parse_args()

    cfg = load_config(args.conf)
    setup_logging(cfg)

    url = args.url or args.url_opt
    if url:
        cfg.set("source", "url", url)

    verify_setting = cfg.get("source", "verify", fallback="true").strip()
    if verify_setting.lower() in {"true", "yes", "1"}:
        verify = True
    elif verify_setting.lower() in {"false", "no", "0"}:
        verify = False
        log.warning("Verificación TLS desactivada en source.verify.")
    else:
        verify = resolve_path(verify_setting)
        if not os.path.isfile(verify):
            log.error(f"Bundle de certificados no encontrado: {verify}")
            sys.exit(2)

    calendar = fetch_calendar(
        cfg.get("source", "url"),
        cfg.getint("source", "timeout", fallback=30),
        verify=verify,
    )
    log.info(
        f"Competición: {calendar.get('competicion')} | {calendar.get('grupo')} "
        f"| Temporada {calendar.get('temporada')}"
    )
    rows = extract_rows(calendar)

    if args.team:
        team = args.team
    elif args.no_prompt:
        team = cfg.get("highlight", "team", fallback="")
    else:
        team = choose_team(list_teams(rows), cfg.get("highlight", "team", fallback=""))
    cfg.set("highlight", "team", team)

    write_excel(rows, calendar, cfg)
    if not args.no_pdf and cfg.getboolean("output", "pdf", fallback=True):
        write_pdf(rows, calendar, cfg)
        compact_rows = [row for row in rows if is_highlighted(row, team.upper())]
        if team and cfg.getboolean("output", "compact_pdf", fallback=True):
            write_pdf(
                compact_rows,
                calendar,
                cfg,
                filename_key="compact_pdf_filename",
                label="PDF compacto",
                highlight_rows=False,
            )
        elif not team:
            log.warning("No se genera el PDF compacto porque no hay equipo seleccionado.")


if __name__ == "__main__":
    try:
        main()
    except requests.RequestException as exc:
        log.error(f"Error de red: {exc}")
        sys.exit(3)
